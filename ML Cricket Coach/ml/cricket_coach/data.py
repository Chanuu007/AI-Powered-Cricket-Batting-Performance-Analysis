"""Read the workbook without modifying it. Preserve literal NA categories."""
from datetime import date, datetime
from pathlib import Path
import re
import numpy as np
import openpyxl
import pandas as pd
from .config import COLUMN_MAP, CONTEXT, DATASET, LABELS, SUPPORTED_FORMATS, LIMITATIONS
from .io_utils import file_hash
from .label_rules import rubric_label


def parse_match_date(value):
    if value is None or pd.isna(value): return pd.NaT
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return pd.Timestamp(value).normalize()
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try: return pd.Timestamp(datetime.strptime(text, fmt)).normalize()
        except ValueError: pass
    return pd.NaT


def _identifier(value):
    if isinstance(value, (int, np.integer)): return str(value)
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    return str(value).strip()


def family_root(match_id, parents):
    """Group related resampled matches without dividing the input into origins."""
    current, seen = str(match_id), set()
    while current in parents:
        if current in seen: raise ValueError("Cycle in source-match notes.")
        seen.add(current)
        current = parents[current]
    return current


def load_dataset(path=DATASET):
    path = Path(path).resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Dataset not found: {path}. Put the workbook in data/Match_Data_Newl.xlsx or pass --dataset.")
    workbook = openpyxl.load_workbook(path, data_only=True)
    if "Match_Data" not in workbook.sheetnames:
        raise ValueError("Workbook needs a worksheet named Match_Data.")
    sheet = workbook["Match_Data"]
    notes, parents = [], {}
    for row in sheet:
        for cell in row:
            if cell.comment:
                text = cell.comment.text
                notes.append({"cell": cell.coordinate, "text": text})
                m = re.search(r"Synthetic match\s+(\d+);\s+resampled source match\s+(\d+)", text, re.I)
                if m: parents[m.group(1)] = m.group(2)
    frame = pd.read_excel(path, sheet_name="Match_Data", keep_default_na=False, engine="openpyxl")
    frame.columns = [str(c).strip() for c in frame.columns]
    missing = [c for c in COLUMN_MAP if c not in frame.columns]
    if missing: raise ValueError(f"Missing worksheet columns: {', '.join(missing)}")
    frame = frame[list(COLUMN_MAP)].rename(columns=COLUMN_MAP).copy()
    frame["record_id"] = [f"row_{r}" for r in range(2, len(frame) + 2)]
    frame["excel_row"] = range(2, len(frame) + 2)
    issues = []
    def issue(index, field, message, old="", new=""):
        issues.append({"record_id": frame.loc[index, "record_id"], "field": field,
                       "issue": message, "original_value": str(old), "processed_value": str(new)})
    for col in ("player_id", "match_id"):
        frame[col] = frame[col].map(_identifier)
        if frame[col].isin(["", "None", "nan"]).any(): raise ValueError(f"Missing {col}.")
    for col in ("runs", "balls_faced", "innings_number"):
        values = pd.to_numeric(frame[col], errors="coerce")
        invalid = values.isna() | ~np.isfinite(values) | (values < 0) | (values % 1 != 0)
        if invalid.any(): raise ValueError(f"Invalid {col} at Excel rows {frame.loc[invalid, 'excel_row'].tolist()[:8]}")
        frame[col] = values.astype(int)
    if (~frame["innings_number"].between(1, 4)).any(): raise ValueError("Innings Number must be between 1 and 4.")
    for col in ["out_status", "dismissal_type", "match_format", "performance_label"] + CONTEXT:
        frame[col] = frame[col].map(lambda x: str(x).strip() or "Unknown")
    if not set(frame["performance_label"]).issubset(LABELS):
        raise ValueError("Performance_Label must contain only Poor, Average, Good, with no blanks.")
    if not set(frame["match_format"]).issubset(SUPPORTED_FORMATS):
        raise ValueError("This project rubric supports T20, ODI and Test only.")
    if not set(frame["out_status"]).issubset(["Out", "Not Out"]):
        raise ValueError("Out / Not Out must be Out or Not Out.")
    invalid = ((frame["out_status"] == "Not Out") != (frame["dismissal_type"] == "Not Out"))
    if invalid.any(): raise ValueError("Dismissal Type conflicts with Out / Not Out.")
    invalid = (frame["match_format"] != "Test") & (frame["innings_number"] > 2)
    if invalid.any(): raise ValueError("A limited-overs record has innings number above 2.")
    if frame.duplicated(["player_id", "match_id", "innings_number"]).any():
        raise ValueError("Duplicate player-match-innings records. Resolve them in the source workbook.")
    supplied_sr = pd.to_numeric(frame["strike_rate"], errors="coerce")
    frame["strike_rate"] = np.divide(100.0 * frame["runs"], frame["balls_faced"],
                                     out=np.zeros(len(frame)), where=frame["balls_faced"].to_numpy() != 0)
    mismatch = supplied_sr.isna() | ((supplied_sr - frame["strike_rate"].round(2)).abs() > 0.011)
    for i in frame.index[mismatch]:
        issue(i, "strike_rate", "Recalculated in processed data only", supplied_sr[i], frame.loc[i, "strike_rate"])
    for i in frame.index[frame["balls_faced"] == 0]:
        issue(i, "strike_rate", "Zero balls faced: encoded SR=0 for the existing rubric", supplied_sr[i], 0)
    frame["match_date"] = frame["match_date"].map(parse_match_date)
    asof = pd.Timestamp.now(tz="Asia/Colombo").tz_localize(None).normalize()
    for i in frame.index[frame["match_date"].isna()]: issue(i, "match_date", "Unparseable; excluded from dated player analysis")
    for i in frame.index[frame["match_date"] > asof]:
        issue(i, "match_date", "Recorded future date; verify possible day/month conversion", frame.loc[i, "match_date"])
    for col in CONTEXT:
        # A run-out has no bowler-attributed dismissal; a not-out has no dismissal delivery.
        na_mask = (frame["out_status"] == "Not Out") | (frame["dismissal_type"] == "Run Out")
        for i in frame.index[na_mask & (frame[col] != "NA")]:
            issue(i, col, "No applicable dismissal delivery; normalized in processed copy", frame.loc[i, col], "NA")
        frame.loc[na_mask, col] = "NA"
    expected = frame.apply(lambda r: rubric_label(r.runs, r.balls_faced, r.match_format), axis=1)
    for i in frame.index[expected != frame["performance_label"]]:
        issue(i, "performance_label", "Stored label differs from documented rubric; retained, not overwritten", frame.loc[i, "performance_label"], expected[i])
    frame["split_group"] = frame["match_id"].map(lambda x: "family_" + family_root(x, parents))
    for col in ("match_format", "match_date"):
        if (frame.groupby("match_id")[col].nunique(dropna=False) > 1).any():
            raise ValueError(f"The same Match ID has conflicting {col} values.")
    players = {}
    if "Players" in workbook.sheetnames:
        for row in list(workbook["Players"].values)[1:]:
            if row[0] is not None: players[_identifier(row[0])] = str(row[1])
    meta = {
        "dataset_filename": path.name, "dataset_sha256": file_hash(path), "rows": len(frame),
        "players": frame.player_id.nunique(), "matches": frame.match_id.nunique(),
        "split_groups": frame.split_group.nunique(), "label_counts": frame.performance_label.value_counts().to_dict(),
        "format_counts": frame.match_format.value_counts().to_dict(), "source_notes": notes,
        "grouping": "Source-match families from existing notes, falling back to Match ID; grouping columns are never model inputs.",
        "date_policy": "Day-first text parsed explicitly; existing Excel date values retained. No inferred date corrections. Verify ambiguous original dates.",
        "input_policy": "All rows processed in one combined table; no origin-based data subsets are created.",
        "limitations": LIMITATIONS,
    }
    workbook.close()
    issue_frame = pd.DataFrame(issues, columns=["record_id", "field", "issue", "original_value", "processed_value"])
    return frame, players, meta, issue_frame
